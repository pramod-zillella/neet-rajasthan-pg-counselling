"""Build filterable Rajasthan 2024 and 2025 seat-matrix tables from PDFs.

2024 values are extracted directly from digest-checked Board PDFs. The 2025
row-level research CSV is reconciled to the four Board PDF snapshots. Round 1
is offered seats; later rounds are separate vacancy snapshots, not additions.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path

import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from fetch_2024_matrix_pdfs import SOURCES, archive_sha1, archived_url


ROOT = Path(__file__).resolve().parents[1]
SITE_BASE = "https://pramod-zillella.github.io/neet-rajasthan-pg-counselling/"
COLUMNS = ["year", "round", "matrix_date", "measure", "route", "institution", "institution_type",
           "branch", "seat_type", "source_seat_label", "seats", "source_id", "pdf_page"]
ROUND_KEYS = {"Round 1": "round-1", "Round 2": "round-2", "Round 3": "round-3",
              "Stray vacancy": "stray", "Stray": "stray", "Special stray vacancy": "special-stray"}
ALIASES = {
    "SP MC Bikaner": "SPMC Bikaner", "SN MC Jodhpur": "Dr SNMC Jodhpur",
    "GMC, Kota": "GMC Kota", "Jhalawar MC": "JMC Jhalawar", "RUHS": "RUHS CMS Jaipur",
    "Bharatpur": "GMC Bharatpur", "Churu": "GMC Churu", "Pali": "GMC Pali",
    "Bhilwara": "GMC Bhilwara", "Dungarpur": "GMC Dungarpur",
    "MG MC Jaipur": "Mahatma Gandhi Medical College and Hospital, Jaipur",
    "NIMS Jaipur": "National Institute of Medical Sciences & Research, (NIMS) Jaipur",
    "Geetanjali MC Udaipur": "Geetanjali Medical College & Hospital, Udaipur",
    "Geetanjali Udaipur": "Geetanjali Medical College & Hospital, Udaipur",
    "Pacific Medical College & Hospital Udaipur": "Pacific Medical College & Hospital, Udaipur",
    "Pacific Medical College & Hospital": "Pacific Medical College & Hospital, Udaipur",
    "Pacific Institute of Medical Sciences Udaipur": "Pacific Institute of Medical Sciences, Umarda, Udaipur",
    "Pacific Institute of Medical Sciences": "Pacific Institute of Medical Sciences, Umarda, Udaipur",
    "JNU Medical College, Jaipur": "Jaipur National University Institute of Medical Sciences and Resarch Centre",
    "American Int. Inst. of Med. Sc., Udaipur": "American International Institute of Medical Sciences, Bedwas",
    "Ananta Institute of Med. Sc., Rajsamand": "Ananta Institute of Medical Sciences & Research Centre, Rajsamand",
    "Dr.SS Tantia Medical College,Hospital and Research Center, Sriganganagar": "Dr S S Tantia Medical College Hospital & Research Centre, Sri Ganganagar",
}


def clean(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def num(value: str | None) -> int:
    value = clean(value)
    assert value.isdigit(), f"Expected a printed seat count, got {value!r}"
    return int(value)


def add(rows: list[dict], *, year: int, round_name: str, date: str, route: str,
        institution: str, institution_type: str, branch: str, seat_type: str,
        source_seat_label: str, seats: int, source_id: str, page: int) -> None:
    rows.append({
        "year": year, "round": round_name, "matrix_date": date,
        "measure": "Initial offered seats" if round_name == "Round 1" else "Published vacant-seat snapshot",
        "route": route, "institution": institution, "institution_type": institution_type,
        "branch": clean(branch), "seat_type": seat_type, "source_seat_label": source_seat_label,
        "seats": seats, "source_id": source_id, "pdf_page": page,
    })


def parse_2024_medical(path: Path, source_id: str, round_name: str, date: str) -> list[dict]:
    rows: list[dict] = []
    with fitz.open(path) as pdf:
        assert len(pdf) == 3, path
        tables = [page.find_tables().tables[0].extract() for page in pdf]
    page_courses = []
    for page_no, table in enumerate(tables, 1):
        headers = table[0]
        detail = [row for row in table if clean(row[0]).isdigit() and clean(row[1])]
        courses = {num(row[0]): clean(row[1]) for row in detail}
        assert len(courses) == len(detail) and len(courses) >= 30, (path, page_no, len(detail))
        page_courses.append(courses)
        if page_no == 1:
            assert len(headers) == 33 and len(table[2]) == 33, path
            for row in detail:
                assert sum(num(v) for v in row[2:28]) == num(row[28]), (path, row[1], "government total")
                for start in range(2, 28, 2):
                    raw_name = clean(headers[start])
                    institution = ALIASES.get(raw_name, raw_name)
                    assert institution != raw_name or raw_name in {"SMS MC Jaipur", "RNT MC Udaipur", "JLN MC Ajmer", "RUHS CMS Jaipur", "GMC Bharatpur", "GMC Churu", "GMC Pali", "GMC Bhilwara", "GMC Dungarpur"}, raw_name
                    for col, seat_type in [(start, "State (Government)"), (start + 1, "State (Management)")]:
                        add(rows, year=2024, round_name=round_name, date=date, route="MD/MS",
                            institution=institution, institution_type="Government", branch=row[1],
                            seat_type=seat_type, source_seat_label=clean(table[2][col]),
                            seats=num(row[col]), source_id=source_id, page=page_no)
                assert sum(num(row[i]) for i in (30, 31, 32)) == num(row[29]), (path, row[1], "MGMC total")
                for col, seat_type in [(30, "AI"), (31, "State"), (32, "Management")]:
                    add(rows, year=2024, round_name=round_name, date=date, route="MD/MS",
                        institution=ALIASES["MG MC Jaipur"], institution_type="Private", branch=row[1],
                        seat_type=seat_type, source_seat_label=clean(table[1][col]),
                        seats=num(row[col]), source_id=source_id, page=page_no)
        else:
            assert len(headers) in {14, 18, 23}, (path, page_no, len(headers))
            for row in detail:
                for start in range(2, len(headers) - 3, 4):
                    raw_name = clean(headers[start])
                    if raw_name.startswith(("Total Private", "Grand Total")):
                        continue
                    institution = ALIASES.get(raw_name)
                    assert institution, (path, page_no, raw_name)
                    assert sum(num(row[i]) for i in (start + 1, start + 2, start + 3)) == num(row[start]), (path, page_no, row[1], raw_name)
                    for offset, seat_type in [(1, "AI"), (2, "State"), (3, "Management")]:
                        col = start + offset
                        add(rows, year=2024, round_name=round_name, date=date, route="MD/MS",
                            institution=institution, institution_type="Private", branch=row[1],
                            seat_type=seat_type, source_seat_label=clean(table[1][col]),
                            seats=num(row[col]), source_id=source_id, page=page_no)
        if page_no == 3 and len(headers) == 23:
            for row in detail:
                from_rows = sum(r["seats"] for r in rows if r["branch"] == clean(row[1]) and r["institution_type"] == "Private")
                government = sum(r["seats"] for r in rows if r["branch"] == clean(row[1]) and r["institution_type"] == "Government")
                assert from_rows == num(row[18]) and government + from_rows == num(row[22]), (path, row[1], "grand total")
    assert page_courses[0] == page_courses[1] == page_courses[2], path
    return rows


def parse_2024_dnb(path: Path, source_id: str, round_name: str, date: str) -> list[dict]:
    rows: list[dict] = []
    with fitz.open(path) as pdf:
        assert len(pdf) == 1, path
        table = pdf[0].find_tables().tables[0].extract()
    assert len(table) == 32 and len(table[0]) == 15, path
    headings = [clean(x) for x in table[0][2:14]]
    for serial, row in enumerate(table[1:-1], 1):
        assert num(row[0]) == serial, (path, serial)
        raw_name = clean(row[1])
        if raw_name.startswith("Diploma) "):
            raw_name = raw_name.removeprefix("Diploma) ")
        if raw_name.endswith("(Post MBBS"):
            raw_name += " Diploma)"
        match = re.fullmatch(r"(.+?) \(Post MBBS (DNB|Diploma)\)", raw_name)
        assert match, (path, serial, raw_name)
        institution = match.group(1)
        route = "Post MBBS DNB" if match.group(2) == "DNB" else "Post MBBS Diploma"
        assert sum(num(v) for v in row[2:14]) == num(row[14]), (path, serial, "row total")
        for branch, value in zip(headings, row[2:14]):
            add(rows, year=2024, round_name=round_name, date=date, route=route,
                institution=institution, institution_type="DNB / Diploma centre", branch=branch,
                seat_type="Programme total", source_seat_label="No quota split in PDF",
                seats=num(value), source_id=source_id, page=1)
    for col in range(2, 15):
        actual = sum(num(row[col]) for row in table[1:-1])
        assert actual == num(table[-1][col]), (path, col, actual, table[-1][col])
    return rows


def matrix_url(row: dict) -> str:
    return f"{SITE_BASE}documents/{row['year']}/matrix/{row['source_id'].removeprefix(str(row['year']) + '-')}.pdf#page={row['pdf_page']}"


def save_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=COLUMNS + ["source_pdf_url"])
        writer.writeheader()
        for row in rows:
            writer.writerow({**row, "source_pdf_url": matrix_url(row)})


def save_excel(by_year: dict[int, list[dict]]) -> None:
    wb = Workbook()
    notes = wb.active
    notes.title = "Read first"
    for note in [
        "Rajasthan NEET-PG seat matrices, 2024 and 2025 academic exam years",
        "Round 1 rows are offered seats; later rounds show vacancies at separate dates. Never add round totals together.",
        "2024 Medical plus DNB/Diploma matrices are included; the 2025 verified source set covers MD/MS medical colleges only.",
        "2024 private AIQ and 2025 AI are labels within Rajasthan Board counselling, not MCC AIQ seats.",
        "Zero-seat rows are retained so filters show which college/course combinations appeared in the source matrix.",
        "Click the Source PDF Page cell for the relevant original Board PDF page.",
    ]:
        notes.append([note])
    notes.column_dimensions["A"].width = 118
    for year in (2024, 2025):
        ws = wb.create_sheet(f"{year} Matrix")
        labels = {"institution": "College / Centre", "branch": "Course / Branch", "seat_type": "Seat Type",
                  "source_seat_label": "Printed Seat Label", "seats": "Seats", "pdf_page": "Source PDF Page"}
        ws.append([labels.get(c, c.replace("_", " ").title()) for c in COLUMNS])
        for i, row in enumerate(by_year[year], 2):
            ws.append([row[c] for c in COLUMNS])
            cell = ws.cell(i, COLUMNS.index("pdf_page") + 1)
            cell.hyperlink = matrix_url(row)
            cell.font = Font(color="0A6B69", underline="single")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="184D55")
        for col, width in {"round": 22, "matrix_date": 35, "measure": 32, "route": 23,
                           "institution": 54, "institution_type": 23, "branch": 45,
                           "seat_type": 25, "source_seat_label": 23}.items():
            ws.column_dimensions[get_column_letter(COLUMNS.index(col) + 1)].width = width
        table = Table(displayName=f"Matrix{year}", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    wb.save(ROOT / "data" / "Rajasthan-NEET-PG-2024-2025-seat-matrices.xlsx")


def build(source_root: Path) -> None:
    by_year: dict[int, list[dict]] = {2024: [], 2025: []}
    documents = []
    for slug, round_name, route_group, date, capture, original, expected_sha1 in SOURCES:
        path = ROOT / "documents" / "2024" / "matrix" / f"{slug}.pdf"
        data = path.read_bytes()
        assert data.startswith(b"%PDF") and archive_sha1(data) == expected_sha1, path
        source_id = f"2024-{slug}"
        rows = (parse_2024_medical(path, source_id, round_name, date)
                if route_group == "Medical" else parse_2024_dnb(path, source_id, round_name, date))
        by_year[2024].extend(rows)
        with fitz.open(path) as pdf:
            pages = len(pdf)
        documents.append({"year": 2024, "round": round_name, "roundKey": ROUND_KEYS[round_name],
                          "routeGroup": "MD/MS" if route_group == "Medical" else "DNB / Diploma",
                          "date": date, "sourceId": source_id, "path": path.relative_to(ROOT).as_posix(),
                          "pages": pages, "rows": len(rows), "seats": sum(r["seats"] for r in rows),
                          "sha256": hashlib.sha256(data).hexdigest(), "archiveSha1": expected_sha1,
                          "sourceUrl": original, "archiveUrl": archived_url(capture, original)})
        print(f"2024 {slug}: {len(rows)} matrix cells, {sum(r['seats'] for r in rows)} seats")

    source_path = source_root / "rajasthan-pg2025"
    with (source_path / "sources.csv").open(encoding="utf-8-sig", newline="") as stream:
        sources25 = {r["Source_ID"]: r for r in csv.DictReader(stream)}
    with (source_path / "seat-matrices.csv").open(encoding="utf-8-sig", newline="") as stream:
        data25 = list(csv.DictReader(stream))
    round_sources = {"Round 1": ("r1-matrix", "round-1"), "Round 2": ("drive-r2-matrix", "round-2"),
                     "Round 3": ("drive-r3-matrix", "round-3"), "Stray": ("stray-matrix", "stray")}
    for round_name, (original_id, slug) in round_sources.items():
        source = sources25[original_id]
        original_path = source_path / source["Local_file"]
        data = original_path.read_bytes()
        assert hashlib.sha256(data).hexdigest() == source["SHA256"], original_path
        target = ROOT / "documents" / "2025" / "matrix" / f"{slug}.pdf"
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(original_path, target)
        source_id = f"2025-{slug}"
        source_rows = [r for r in data25 if r["Round"] == round_name]
        assert len(source_rows) == 1536, (round_name, len(source_rows))
        rows = []
        for row in source_rows:
            assert row["Measure"] == ("Initial offered seats" if round_name == "Round 1" else "Published vacant-seat snapshot")
            assert row["Source_ID"] == original_id
            page = num(row["PDF_page"])
            assert 1 <= page <= int(source["PDF_pages"]), row
            add(rows, year=2025, round_name=round_name, date=row["Matrix_date"], route="MD/MS",
                institution=clean(row["Institution"]), institution_type=clean(row["Institution_type"]),
                branch=row["Branch"], seat_type=clean(row["Seat_type"]),
                source_seat_label=clean(row["Seat_type"]), seats=num(row["Seats"]),
                source_id=source_id, page=page)
        by_year[2025].extend(rows)
        with fitz.open(target) as pdf:
            assert len(pdf) == int(source["PDF_pages"]), target
        documents.append({"year": 2025, "round": round_name, "roundKey": ROUND_KEYS[round_name],
                          "routeGroup": "MD/MS", "date": source["Document_date"],
                          "sourceId": source_id, "path": target.relative_to(ROOT).as_posix(),
                          "pages": int(source["PDF_pages"]), "rows": len(rows),
                          "seats": sum(r["seats"] for r in rows), "sha256": source["SHA256"],
                          "sourceUrl": source["URL"], "archiveUrl": None})
        print(f"2025 {slug}: {len(rows)} matrix cells, {sum(r['seats'] for r in rows)} seats")
    assert {round_name: sum(r["seats"] for r in by_year[2025] if r["round"] == round_name)
            for round_name in round_sources} == {"Round 1": 2378, "Round 2": 1311, "Round 3": 1122, "Stray": 550}
    keys = [(r["year"], r["round"], r["route"], r["institution"], r["branch"], r["seat_type"]) for rows in by_year.values() for r in rows]
    assert len(keys) == len(set(keys)), [item for item, count in Counter(keys).items() if count > 1][:8]
    for year, rows in by_year.items():
        (ROOT / "data" / f"{year}-matrix.json").write_text(json.dumps(rows, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
        (ROOT / "data" / f"{year}-matrix-available.json").write_text(
            json.dumps([row for row in rows if row["seats"] > 0], ensure_ascii=False, separators=(",", ":")) + "\n",
            encoding="utf-8",
        )
        save_csv(ROOT / "data" / f"{year}-matrix.csv", rows)
    (ROOT / "data" / "matrix-documents.json").write_text(json.dumps({"documents": documents}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    save_excel(by_year)
    print(f"Saved {len(by_year[2024])} 2024 rows and {len(by_year[2025])} 2025 rows")


if __name__ == "__main__":
    build(ROOT.parent / "neet-specialty-knowledge-base" / "exports")
