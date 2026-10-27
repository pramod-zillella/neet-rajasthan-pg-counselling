"""Extract checked, privacy-limited Rajasthan merit/allotment tables.

The Board PDFs are the source of the 2024 rows. The separately validated 2025
research CSVs are the source of the 2025 rows. Names, NEET roll numbers, gender
and relatives' names are intentionally omitted from the published tables.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from pathlib import Path

import fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SITE_BASE = "https://pramod-zillella.github.io/neet-rajasthan-pg-counselling/"
MERIT_FILES = {
    "Round 1": ("round-1.pdf", 4561, (485, 525), (438, 480), (397, 430), (240, 310), (120, 160), (910, 960)),
    "Round 2": ("round-2.pdf", 4684, (485, 525), (438, 480), (397, 430), (240, 310), (120, 160), (910, 960)),
    "Round 3": ("round-3.pdf", 6871, (463, 510), (424, 462), (386, 417), (235, 305), (132, 161), (920, 965)),
    "Stray vacancy": ("stray.pdf", 533, (455, 495), (402, 439), (373, 397), (240, 280), (90, 124), (920, 965)),
    "Special stray vacancy": ("special-stray.pdf", 265, (486, 519), (442, 481), (406, 430), (275, 309), (124, 149), (920, 965)),
}
ALLOTMENT_FILES = {
    "Round 1": ("round-1.pdf", 1506),
    "Round 2": ("round-2.pdf", 1465),
    "Round 3": ("round-3.pdf", 1998),
}
MERIT_2025_BANDS = {
    "Round 1": ("round-1.pdf", 3706, (272, 330), (330, 378), (378, 445)),
    "Round 2": ("round-2.pdf", 3830, (262, 323), (323, 370), (370, 441)),
    "Round 3": ("round-3.pdf", 5289, (252, 289), (289, 332), (332, 416)),
    "Stray": ("stray.pdf", 396, (260, 312), (312, 352), (352, 391)),
}
MERIT_COLUMNS = [
    "year", "round", "list_serial", "registration_id", "state_merit_rank",
    "neet_pg_air", "domicile", "filled_category", "considered_category",
    "service_status", "neet_marks", "bonus_marks", "total_marks", "pdf_page",
]
ALLOTMENT_COLUMNS = [
    "year", "round", "list_serial", "registration_id", "merit_number",
    "neet_pg_air", "college", "course", "seat_type", "domicile", "considered_category",
    "service_status", "allotment_status", "pdf_page",
]
EXCEL_LABELS = {
    "neet_pg_air": "NEET-PG AIR",
    "seat_type": "Allotted Seat Type",
    "considered_category": "Provisional Considered Category",
    "service_status": "Service / Non-Service",
    "pdf_page": "Source PDF Page",
}


def clean(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def integer(value: str | int | None) -> int | None:
    value = clean(str(value)) if value is not None else ""
    return int(value) if value.isdigit() else None


def number(value: str | None) -> int | float | None:
    try:
        parsed = float(value) if value else None
        return int(parsed) if parsed is not None and parsed.is_integer() else parsed
    except ValueError:
        return None


def field(words: list[tuple], band: tuple[int, int]) -> str:
    selected = [w for w in words if band[0] <= w[1] < band[1]]
    return clean(" ".join(w[4] for w in sorted(selected, key=lambda w: -w[1])))


def extract_2024_merit() -> list[dict]:
    result = []
    for round_name, (filename, expected, dom, filled, considered, service, rank, serial) in MERIT_FILES.items():
        pdf = ROOT / "documents" / "2024" / "merit" / filename
        rows = []
        with fitz.open(pdf) as document:
            for page_number, page in enumerate(document, 1):
                words = page.get_text("words")
                anchors = [w for w in words if re.fullmatch(r"PM\d+", w[4])]
                for anchor in anchors:
                    same_row = [w for w in words if abs(w[0] - anchor[0]) < 5]
                    row = {
                        "year": 2024,
                        "round": round_name,
                        "list_serial": integer(field(same_row, serial)),
                        "registration_id": anchor[4],
                        "state_merit_rank": integer(field(same_row, rank)),
                        "neet_pg_air": None,
                        "domicile": field(same_row, dom),
                        "filled_category": field(same_row, filled),
                        "considered_category": field(same_row, considered),
                        "service_status": field(same_row, service),
                        "neet_marks": None,
                        "bonus_marks": None,
                        "total_marks": None,
                        "pdf_page": page_number,
                    }
                    rows.append(row)
        assert len(rows) == expected, (round_name, len(rows), expected)
        serials = [row["list_serial"] for row in rows]
        assert serials == list(range(1, expected + 1)), (round_name, "serial sequence", [(i + 1, v) for i, v in enumerate(serials) if v != i + 1][:6])
        assert all(row["state_merit_rank"] is not None for row in rows), (round_name, "missing rank")
        assert all(row["domicile"] in {"RAJASTHAN", "OTHER"} for row in rows), (round_name, "domicile", Counter(row["domicile"] for row in rows).most_common(5))
        assert all(row["service_status"] in {"In-service", "Non-service", "Sr. Demonstrator", ""} for row in rows), (round_name, "service", Counter(row["service_status"] for row in rows).most_common(8))
        assert all(row["considered_category"] for row in rows), (round_name, "category")
        result.extend(rows)
        print(f"2024 merit {round_name}: {len(rows)} rows")
    return result


def extract_2024_allotments() -> list[dict]:
    result = []
    for round_name, (filename, expected) in ALLOTMENT_FILES.items():
        pdf = ROOT / "documents" / "2024" / "allotment" / filename
        rows = []
        with fitz.open(pdf) as document:
            for page_number, page in enumerate(document, 1):
                words = page.get_text("words")
                anchors = sorted([w for w in words if re.fullmatch(r"PM\d+", w[4])], key=lambda w: w[1])
                for index, anchor in enumerate(anchors):
                    next_y = anchors[index + 1][1] - 3 if index + 1 < len(anchors) else min(page.rect.height - 35, anchor[1] + 18)
                    row_words = [w for w in words if anchor[1] - 4 <= w[1] < next_y]

                    def column(low: float, high: float) -> str:
                        subset = sorted((w for w in row_words if low <= w[0] < high), key=lambda w: (round(w[1], 1), w[0]))
                        return clean(" ".join(w[4] for w in subset))

                    college_raw = column(655, 1008)
                    seat_match = re.search(r"\((State|All India|Mgmt|Post MBBS DNB|Post MBBS Diploma)\s*(?:-\s*Recog)?\)\s*$", college_raw, re.I)
                    seat_type = seat_match.group(1) if seat_match else ""
                    if seat_type == "Mgmt":
                        seat_type = "Management"
                    college = clean(college_raw[:seat_match.start()]) if seat_match else college_raw
                    rows.append({
                        "year": 2024,
                        "round": round_name,
                        "list_serial": integer(column(45, 88)),
                        "registration_id": anchor[4],
                        "merit_number": integer(column(88, 140)),
                        "neet_pg_air": None,
                        "college": college,
                        "course": column(545, 655),
                        "seat_type": seat_type,
                        "domicile": "",
                        "considered_category": "",
                        "service_status": "",
                        "allotment_status": "",
                        "pdf_page": page_number,
                    })
        assert len(rows) == expected, (round_name, len(rows), expected)
        serials = [row["list_serial"] for row in rows]
        assert serials == list(range(1, expected + 1)), (round_name, "serial sequence", [(i + 1, v) for i, v in enumerate(serials) if v != i + 1][:6])
        assert all(row["merit_number"] is not None for row in rows), (round_name, "missing rank")
        assert all(row["college"] and row["course"] for row in rows), (round_name, "college/course missing", [r for r in rows if not r["college"] or not r["course"]][:3])
        assert all(row["seat_type"] for row in rows), (round_name, "seat type missing", [r for r in rows if not r["seat_type"]][:3])
        result.extend(rows)
        print(f"2024 allotment {round_name}: {len(rows)} rows; seat labels {Counter(r['seat_type'] for r in rows)}")
    return result


def enrich_2024_allotments(allot: list[dict], merit: list[dict]) -> None:
    by_id = {(r["round"], r["registration_id"]): r for r in merit}
    assert len(by_id) == len(merit), "Duplicate registration ID within a merit round"
    for row in allot:
        match = by_id.get((row["round"], row["registration_id"]))
        assert match, (row["round"], row["registration_id"])
        assert match["state_merit_rank"] == row["merit_number"], (row, match)
        row["domicile"] = match["domicile"]
        row["considered_category"] = match["considered_category"]
        row["service_status"] = match["service_status"]


def extract_2025_merit_details() -> dict[tuple[str, str], dict[str, str]]:
    details = {}
    for round_name, (filename, expected, dom_band, filled_band, considered_band) in MERIT_2025_BANDS.items():
        pdf = ROOT / "documents" / "2025" / "merit" / filename
        count = 0
        with fitz.open(pdf) as document:
            for page_number, page in enumerate(document, 1):
                words = page.get_text("words")
                for anchor in (w for w in words if re.fullmatch(r"PM\d+\**", w[4])):
                    row_words = [w for w in words if abs(w[1] - anchor[1]) < 4]

                    def xfield(band: tuple[int, int]) -> str:
                        return clean(" ".join(w[4] for w in sorted(row_words, key=lambda w: w[0]) if band[0] <= w[0] < band[1]))

                    reg_id = anchor[4].rstrip("*")
                    key = (round_name, reg_id)
                    assert key not in details, key
                    details[key] = {
                        "domicile": xfield(dom_band),
                        "filled_category": xfield(filled_band),
                        "considered_category": xfield(considered_band),
                        "pdf_page": page_number,
                    }
                    count += 1
        assert count == expected, (round_name, count, expected)
        print(f"2025 merit details {round_name}: {count} rows")
    assert all(r["domicile"] for r in details.values()), Counter(r["domicile"] for r in details.values()).most_common(20)
    assert all(r["considered_category"] for r in details.values()), "Missing 2025 provisional considered category"
    return details


def normalize_2025_seat_and_course(raw_seat: str, raw_course: str) -> tuple[str, str]:
    seat = clean(raw_seat)
    course = clean(raw_course)
    if seat.startswith("Lepros"):
        seat = seat.removeprefix("Lepros")
    elif not seat and course.endswith("LeprosAI"):
        seat = "AI"
        course = course.removesuffix("AI")
    if "Skin & V.D" in course or "Dermatology Ven." in course:
        # The source uses several variants. In R3, the course text also
        # collides with the seat-type column. Use an explicit normalized
        # specialty label rather than publishing the corrupted extraction.
        course = "MD (Dermatology, Venereology and Leprosy)"
    assert seat in {"Government", "AI", "Management"}, (raw_seat, raw_course)
    return seat, course


def normalize_2025_domiciles(merit: list[dict]) -> None:
    """Repair state names clipped into the adjacent category column in some PDFs.

    The long-name repairs were checked against another round for the same
    registration ID where available. The PDF only prints a shortened Dadra
    and Nagar label, so its unpublished ending is deliberately not inferred.
    """
    repairs = {
        "Madhya PradeGeneral": "Madhya Pradesh",
        "Madhya PradeOBC": "Madhya Pradesh",
        "Andhra PradesGeneral": "Andhra Pradesh",
        "Andhra PradesOBC": "Andhra Pradesh",
        "Himachal PradGeneral": "Himachal Pradesh",
        "Jammu and KaGeneral": "Jammu and Kashmir",
        "Jammu and KaST": "Jammu and Kashmir",
        "Jammu and KashmGeneral": "Jammu and Kashmir",
        "Dadra and NagGeneral": "Dadra and Nagar… (PDF truncated)",
        "Dadra and Nagar": "Dadra and Nagar… (PDF truncated)",
    }
    by_id: dict[str, set[str]] = {}
    for row in merit:
        if row["domicile"] not in repairs and row["domicile"] != "General":
            by_id.setdefault(row["registration_id"], set()).add(row["domicile"])
    for row in merit:
        raw = row["domicile"]
        if raw == "General":
            alternatives = by_id.get(row["registration_id"], set())
            assert len(alternatives) == 1, (row["registration_id"], alternatives)
            row["domicile"] = next(iter(alternatives))
        else:
            row["domicile"] = repairs.get(raw, raw)
    assert not any("General" in row["domicile"] or "OBC" in row["domicile"] for row in merit)


def load_2025(source_root: Path) -> tuple[list[dict], list[dict]]:
    folder = source_root / "rajasthan-pg2025" / "merit-and-air"
    with (folder / "merit-all-rounds.csv").open(encoding="utf-8-sig", newline="") as stream:
        merit_raw = list(csv.DictReader(stream))
    with (folder / "allotments-with-AIR.csv").open(encoding="utf-8-sig", newline="") as stream:
        allot_raw = list(csv.DictReader(stream))
    detail = extract_2025_merit_details()
    merit = [{
        "year": 2025, "round": clean(r["Round"]), "list_serial": None,
        "registration_id": clean(r["Registration_ID"]),
        "state_merit_rank": integer(r["State_merit_rank"]),
        "neet_pg_air": integer(r["NEET_PG_AIR"]),
        "domicile": detail[(r["Round"], r["Registration_ID"])]["domicile"],
        "filled_category": detail[(r["Round"], r["Registration_ID"])]["filled_category"],
        "considered_category": detail[(r["Round"], r["Registration_ID"])]["considered_category"],
        "service_status": clean(r["Service_status"]),
        "neet_marks": number(r["NEET_marks"]),
        "bonus_marks": number(r["Bonus_marks"]),
        "total_marks": number(r["Total_marks"]),
        "pdf_page": integer(r["Merit_PDF_page"]),
    } for r in merit_raw]
    assert all(detail[(r["round"], r["registration_id"])]["pdf_page"] == r["pdf_page"] for r in merit)
    normalize_2025_domiciles(merit)
    merit_by_id = {(r["round"], r["registration_id"]): r for r in merit}
    assert len(merit_by_id) == len(merit)
    allot = [{
        "year": 2025, "round": clean(r["Round"]), "list_serial": None,
        "registration_id": clean(r["Registration_ID"]),
        "merit_number": integer(r["Allotment_state_merit"]),
        "neet_pg_air": integer(r["NEET_PG_AIR"]),
        "college": clean(r["College"]),
        "course": normalize_2025_seat_and_course(r["Seat_type"], r["Course"])[1],
        "seat_type": normalize_2025_seat_and_course(r["Seat_type"], r["Course"])[0],
        "domicile": merit_by_id[(r["Round"], r["Registration_ID"])]["domicile"],
        "considered_category": merit_by_id[(r["Round"], r["Registration_ID"])]["considered_category"],
        "service_status": clean(r["Service_status"]),
        "allotment_status": clean(r["Allotment_status"]),
        "pdf_page": integer(r["Allotment_PDF_page"]),
    } for r in allot_raw]
    assert len(merit) == 3706 + 3830 + 5289 + 396
    assert len(allot) == 1604 + 1551 + 2145 + 220
    assert all(r["neet_pg_air"] and r["state_merit_rank"] for r in merit)
    assert all(r["neet_pg_air"] and r["merit_number"] and r["college"] and r["course"] for r in allot)
    assert all(merit_by_id[(r["round"], r["registration_id"])]["state_merit_rank"] == r["merit_number"] for r in allot)
    assert all(merit_by_id[(r["round"], r["registration_id"])]["neet_pg_air"] == r["neet_pg_air"] for r in allot)
    return merit, allot


def save_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns + ["source_pdf_url"])
        writer.writeheader()
        kind = "merit" if "state_merit_rank" in columns else "allotment"
        for row in rows:
            writer.writerow({**row, "source_pdf_url": pdf_url(row, kind)})


def pdf_url(row: dict, kind: str) -> str:
    label = row["round"].lower()
    slug = "special-stray" if "special" in label else "stray" if "stray" in label else "round-1" if "round 1" in label else "round-2" if "round 2" in label else "round-3"
    return f"{SITE_BASE}documents/{row['year']}/{kind}/{slug}.pdf#page={row['pdf_page']}"


def save_workbook(data: dict[tuple[int, str], list[dict]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for year, kind in [(2024, "merit"), (2024, "allotment"), (2025, "merit"), (2025, "allotment")]:
        rows = data[(year, kind)]
        columns = MERIT_COLUMNS if kind == "merit" else ALLOTMENT_COLUMNS
        ws = workbook.create_sheet(f"{year} {'Merit' if kind == 'merit' else 'Allotments'}")
        ws.append([EXCEL_LABELS.get(column, column.replace("_", " ").title()) for column in columns])
        for excel_row, row in enumerate(rows, 2):
            ws.append([row.get(column) for column in columns])
            source_cell = ws.cell(excel_row, columns.index("pdf_page") + 1)
            source_cell.hyperlink = pdf_url(row, kind)
            source_cell.font = Font(color="0A6B69", underline="single")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="184D55")
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 23
        widths = {"round": 21, "registration_id": 20, "college": 48, "course": 34, "allotment_status": 27, "filled_category": 22,
                  "considered_category": 22, "service_status": 18, "domicile": 18}
        for idx, column in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(idx)].width = widths.get(column, 17)
        table = Table(displayName=f"T{year}{kind.title()}", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    notes = workbook.create_sheet("Read first", 0)
    for line in [
        "Rajasthan NEET-PG 2024 and 2025 · historical merit and allotment tables",
        "2024 merit lists do not show NEET-PG All India Rank; blank AIR cells are intentional.",
        "2024 Round 1 allotment Merit Number is combined merit; later 2024 rounds use state merit after weightage.",
        "2024 Stray and Special Stray allotments are archive links only and have no extracted rows.",
        "The recovered 2025 Round 3 allotment PDF is incomplete or unverified (74 physical pages; footer says of 125).",
        "A few 2025 domicile names ran into the category column; clear state prefixes were normalized. Dadra and Nagar remains marked as truncated in the PDF.",
        "Provisional allotments do not prove final joining. Candidates can recur across rounds; counts are not unique candidates.",
        "Names, relatives, gender and NEET roll identifiers are omitted from these derived tables.",
        "Click each PDF Page cell to open the corresponding physical PDF page on the website.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 115
    notes["A1"].font = Font(bold=True, size=14, color="184D55")
    workbook.save(ROOT / "data" / "Rajasthan-NEET-PG-2024-2025-filterable.xlsx")


def main(source_root: Path) -> None:
    data = {
        (2024, "merit"): extract_2024_merit(),
        (2024, "allotment"): extract_2024_allotments(),
    }
    enrich_2024_allotments(data[(2024, "allotment")], data[(2024, "merit")])
    data[(2025, "merit")], data[(2025, "allotment")] = load_2025(source_root)
    for (year, kind), rows in data.items():
        basename = f"{year}-{kind}"
        (ROOT / "data" / f"{basename}.json").write_text(
            json.dumps(rows, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8"
        )
        save_csv(ROOT / "data" / f"{basename}.csv", rows, MERIT_COLUMNS if kind == "merit" else ALLOTMENT_COLUMNS)
        print(f"Saved {basename}: {len(rows)} rows")
    save_workbook(data)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT.parent / "neet-specialty-knowledge-base" / "exports")
    args = parser.parse_args()
    main(args.source.resolve())
