"""Check public document integrity and local site links before publishing."""

from __future__ import annotations

import hashlib
import json
import csv
import base64
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

from labels import VALID_2025_COLLEGES, VALID_2025_COURSES
from build_matrix import FULL_2025_BRANCH, TRUNCATED_2025_BRANCH


ROOT = Path(__file__).resolve().parents[1]


class Links(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.paths: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_dict = dict(attrs)
        target = attrs_dict.get("href") or attrs_dict.get("src")
        if target and not urlsplit(target).scheme and not target.startswith(("#", "//")):
            self.paths.append(urlsplit(target).path)


def main() -> None:
    manifest = json.loads((ROOT / "data" / "documents.json").read_text(encoding="utf-8"))
    docs = manifest["documents"]
    assert len(docs) == 18, f"Expected 18 round records, got {len(docs)}"
    assert sum(d["kind"] == "merit" for d in docs) == 9
    assert sum(d["kind"] == "allotment" and bool(d["path"]) for d in docs) == 7
    assert sum(d["status"] == "archive-only" for d in docs) == 2
    local_pdfs = set()
    for doc in docs:
        assert doc["sourceUrl"].startswith("http"), doc
        if not doc["path"]:
            assert doc["archiveUrl"].startswith("https://web.archive.org/"), doc
            continue
        path = ROOT / doc["path"]
        assert path.is_file(), path
        assert path.read_bytes().startswith(b"%PDF"), path
        assert hashlib.sha256(path.read_bytes()).hexdigest().lower() == doc["sha256"].lower(), path
        assert path.stat().st_size < 100 * 1024 * 1024, path
        local_pdfs.add(path)
        if doc.get("alternatePath"):
            alternate = ROOT / doc["alternatePath"]
            assert alternate.is_file(), alternate
            local_pdfs.add(alternate)

    for html in ROOT.glob("*.html"):
        parser = Links()
        parser.feed(html.read_text(encoding="utf-8"))
        for relative in parser.paths:
            assert (html.parent / relative).is_file(), f"Broken local link: {html.name} -> {relative}"

    assert len(local_pdfs) == 17
    expected = {(2024,"merit"):16914, (2024,"allotment"):4969,
                (2025,"merit"):13221, (2025,"allotment"):5520}
    rows_by_key = {}
    for (year, kind), count in expected.items():
        stem = f"{year}-{kind}"
        rows = json.loads((ROOT / "data" / f"{stem}.json").read_text(encoding="utf-8"))
        assert len(rows) == count, stem
        assert all(row["year"] == year and row["domicile"] and row["considered_category"] for row in rows), stem
        assert all(row["pdf_page"] > 0 for row in rows), stem
        if kind == "allotment":
            assert all(row["college"] and row["course"] and row["seat_type"] and row["service_status"] for row in rows), stem
            if year == 2025:
                assert {row["college"] for row in rows} == VALID_2025_COLLEGES, "Unreviewed/missing 2025 college label"
                assert {row["course"] for row in rows} == VALID_2025_COURSES, "Unreviewed/missing 2025 course label"
            else:
                assert all(row["course"].startswith(("MD (", "MS (", "Post MBBS ", "Post Diploma ")) for row in rows)
        if year == 2025:
            assert all("General" not in row["domicile"] and "OBC" not in row["domicile"] for row in rows), stem
        assert all(row["neet_pg_air"] is None for row in rows) if year == 2024 else all(row["neet_pg_air"] for row in rows), stem
        with (ROOT / "data" / f"{stem}.csv").open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            assert "source_pdf_url" in reader.fieldnames
            for index, exported in enumerate(reader):
                if kind == "allotment":
                    assert (exported["college"], exported["course"]) == (rows[index]["college"], rows[index]["course"])
            assert index + 1 == count
        rows_by_key[(year,kind)] = rows

    for year in (2024, 2025):
        merit = {(r["round"], r["registration_id"]): r for r in rows_by_key[(year,"merit")]}
        for allot in rows_by_key[(year,"allotment")]:
            source = merit[(allot["round"], allot["registration_id"])]
            assert allot["merit_number"] == source["state_merit_rank"]
            assert allot["neet_pg_air"] == source["neet_pg_air"]
            assert allot["domicile"] == source["domicile"]
            assert allot["considered_category"] == source["considered_category"]
            assert allot["service_status"] == source["service_status"]

    checked_2025 = {(r["round"], r["registration_id"]): r for r in rows_by_key[(2025, "allotment")]}
    for registration, college, course in [
        ("PM0590", "Mahatma Gandhi Medical College and Hospital, Jaipur", "MS (Ophthalmology)"),
        ("PM0129", "Jaipur National University Institute for Medical Sciences and Research Centre, Jaipur.", "MD (General Medicine)"),
        ("PM2491", "Ananta Institute of Medical Sciences and Research Centre, Rajsamand", "MD (Paediatrics)"),
        ("PM2335", "Mahatma Gandhi Medical College and Hospital, Jaipur", "MD (Nuclear Medicine)"),
        ("PM1642", "American International Institute Of Medical Sciences", "MS (General Surgery)"),
        ("PM5283", "Geetanjali Medical College & Hospital, Udaipur", "MS (General Surgery)"),
    ]:
        row = checked_2025[("Round 1" if registration in {"PM0590", "PM0129"} else "Round 3", registration)]
        assert (row["college"], row["course"]) == (college, course), registration

    workbook = load_workbook(ROOT / "data" / "Rajasthan-NEET-PG-2024-2025-filterable.xlsx", read_only=False)
    for (year, kind), count in expected.items():
        sheet = workbook[f"{year} {'Merit' if kind == 'merit' else 'Allotments'}"]
        assert sheet.max_row == count + 1
        headers = [cell.value for cell in sheet[1]]
        assert "Provisional Considered Category" in headers and "Service / Non-Service" in headers
        if kind == "allotment":
            assert "Allotted Seat Type" in headers
            if year == 2025:
                college_col, course_col = headers.index("College") + 1, headers.index("Course") + 1
                for index, row in enumerate(rows, 2):
                    assert (sheet.cell(index, college_col).value, sheet.cell(index, course_col).value) == (row["college"], row["course"])
        page_col = headers.index("Source PDF Page") + 1
        assert sheet.cell(2, page_col).hyperlink.target.startswith("https://pramod-zillella.github.io/")

    matrix_docs = json.loads((ROOT / "data" / "matrix-documents.json").read_text(encoding="utf-8"))["documents"]
    assert len(matrix_docs) == 14
    by_source = {}
    for doc in matrix_docs:
        path = ROOT / doc["path"]
        data = path.read_bytes()
        assert data.startswith(b"%PDF") and hashlib.sha256(data).hexdigest() == doc["sha256"], path
        assert doc["pages"] > 0 and doc["rows"] > 0 and doc["sourceUrl"].startswith("http"), doc
        if doc["year"] == 2024:
            assert base64.b32encode(hashlib.sha1(data).digest()).decode("ascii").rstrip("=") == doc["archiveSha1"], path
        by_source[doc["sourceId"]] = doc
    matrix_book = load_workbook(ROOT / "data" / "Rajasthan-NEET-PG-2024-2025-seat-matrices.xlsx", read_only=False)
    matrix_counts = {2024: 10920, 2025: 6144}
    for year, expected_count in matrix_counts.items():
        rows = json.loads((ROOT / "data" / f"{year}-matrix.json").read_text(encoding="utf-8"))
        assert len(rows) == expected_count and all(r["year"] == year and r["seats"] >= 0 for r in rows)
        if year == 2025:
            assert FULL_2025_BRANCH in {r["branch"] for r in rows}
            assert TRUNCATED_2025_BRANCH not in {r["branch"] for r in rows}
        available = json.loads((ROOT / "data" / f"{year}-matrix-available.json").read_text(encoding="utf-8"))
        assert available == [r for r in rows if r["seats"] > 0]
        assert all(r["source_id"] in by_source and 1 <= r["pdf_page"] <= by_source[r["source_id"]]["pages"] for r in rows)
        assert all(r["measure"] == ("Initial offered seats" if r["round"] == "Round 1" else "Published vacant-seat snapshot") for r in rows)
        for doc in (d for d in matrix_docs if d["year"] == year):
            matched = [r for r in rows if r["source_id"] == doc["sourceId"]]
            assert len(matched) == doc["rows"] and sum(r["seats"] for r in matched) == doc["seats"], doc
        with (ROOT / "data" / f"{year}-matrix.csv").open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            assert "source_pdf_url" in reader.fieldnames
            for index, exported in enumerate(reader):
                source = by_source[rows[index]["source_id"]]
                assert exported["source_pdf_url"] == f"https://pramod-zillella.github.io/neet-rajasthan-pg-counselling/{source['path']}#page={rows[index]['pdf_page']}"
                assert exported["branch"] == rows[index]["branch"]
            assert index + 1 == expected_count
        sheet = matrix_book[f"{year} Matrix"]
        assert sheet.max_row == expected_count + 1
        if year == 2025:
            branch_col = [c.value for c in sheet[1]].index("Course / Branch") + 1
            for index, row in enumerate(rows, 2):
                assert sheet.cell(index, branch_col).value == row["branch"]
        page_col = [c.value for c in sheet[1]].index("Source PDF Page") + 1
        assert sheet.cell(2, page_col).hyperlink.target.startswith("https://pramod-zillella.github.io/")
    assert len(list((ROOT / "documents" / "2024" / "matrix").glob("*.pdf"))) == 10
    assert len(list((ROOT / "documents" / "2025" / "matrix").glob("*.pdf"))) == 4
    print("Verified 40,624 merit/allotment rows, 17,064 matrix cells, round/AIR joins, Excel links, 31 PDFs and site links.")


if __name__ == "__main__":
    main()
