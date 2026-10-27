"""Compare Round 1 MD/MS seats without treating old private pages as 2026 data."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
BASE_URL = "https://pramod-zillella.github.io/neet-rajasthan-pg-counselling/"
SOURCE_2026 = "https://pg.neetrajasthan.com/notifications/1792388159259.pdf"
PDF_2026 = ROOT / "documents/2026/matrix/round-1-medical.pdf"
PDF_2026_SHA256 = "95f9bb153604f332b8394ed14e2b27046077a4a99736adea6425652bc3da2e09"

# The 2024 PDF uses short names; 2025 and the 2026 government page use
# expanded programme labels. Keep this mapping explicit and fail on new labels.
BRANCH_2024 = {
    "MD (Anaesthesia)": "MD (ANAESTHESIOLOGY)",
    "MD (Anatomy)": "MD/MS (ANATOMY)",
    "MS (Anatomy)": "MD/MS (ANATOMY)",
    "MD (Biochemistry)": "MD (BIOCHEMISTRY)",
    "MD (Emergency Medicine)": "MD (EMERGENCY MEDICINE)",
    "MD (Family Medicine)": "MD (FAMILY MEDICINE)",
    "MD (Forensic Medicine)": "MD (FORENSIC MEDICINE)",
    "MD (General Medicine)": "MD (GENERAL MEDICINE)",
    "MD (Geriatric Medicine)": "MD (GERIATRICS/GERIATRIC MEDICINE)",
    "MD (Hospital Administration)": "MD (HOSPITAL ADMINISTRATION)",
    "MD (Immunohaematology and Blood Transfusion)": "MD (IMMUNOHAEMATOLOGY & BLOOD TRANSFUSION/ TRANSFUSION MEDICINE)",
    "MD (Transfusion Medicine)": "MD (IMMUNOHAEMATOLOGY & BLOOD TRANSFUSION/ TRANSFUSION MEDICINE)",
    "MD (Laboratory Medicine)": "MD (LABORATORY MEDICINE)",
    "MD (Microbiology)": "MD (MICROBIOLOGY)",
    "MD (Nuclear Medicine)": "MD (NUCLEAR MEDICINE)",
    "MD (PSM)": "MD (PREVENTIVE & SOCIAL MEDICINE (PSM)/ COMMUNITY MEDICINE)",
    "MD (Paediatrics)": "MD (PAEDIATRICS)",
    "MD (Palliative Medicine)": "MD (PALLIATIVE MEDICINE)",
    "MD (Pathology)": "MD (PATHOLOGY)",
    "MD (Pharmacology)": "MD (PHARMACOLOGY)",
    "MD (Phy. Med. & Reh.)": "MD (PHYSICAL MED. & REHABILITATION)",
    "MD (Physiology)": "MD (PHYSIOLOGY)",
    "MD (Psychiatry)": "MD (PSYCHIATRY)",
    "MD (Radiodiagnosis)": "MD (RADIO-DIAGNOSIS/ RADIOLOGY)",
    "MD (Radiotherapy)": "MD (RADIOTHERAPY/ RADIATION ONCOLOGY)",
    "MD (Skin & V.D)": "MD (SKIN & V.D./DERMATOLOGY VEN. & LEPROSY)",
    "MD (T.B. & Res.Medicine)": "MD (RESPIRATORY MEDICINE/T.B. & CHEST/ TUBERCULOSIS & RESPIRATORY DISEASES)",
    "MD (Tropical Medicine)": "MD (TROPICAL MEDICINE)",
    "MS (ENT)": "MS (E.N.T./ OTORHINOLARYNGOLOGY)",
    "MS (General Surgery)": "MS (GENERAL SURGERY)",
    "MS (Obst. & Gynaecology)": "MS (OBSTETRICS & GYNAECOLOGY)",
    "MS (Ophthalmology)": "MS (OPHTHALMOLOGY)",
    "MS (Orthopaedics)": "MS (ORTHOPAEDICS)",
    "MS (Traumatology & Surgery)": "MS (TRAUMATOLOGY & SURGERY)",
}

INSTITUTION_2026 = {
    "Dr SNMC Jodhpur": "Dr SNMC Jodhpur",
    "GMC Kota": "GMC Kota", "JLN MC Ajmer": "JLN MC Ajmer",
    "RNT MC Udaipur": "RNT MC Udaipur", "SMS MC Jaipur": "SMS MC Jaipur",
    "SPMC Bikaner": "SPMC Bikaner", "ESIC Alwar": "ESIC Alwar",
    "RUHS CMS Jaipur": "RUHS CMS Jaipur", "GMC Bharatpur": "GMC Bharatpur",
    "GMC Bhilwara": "GMC Bhilwara", "GMC Churu": "GMC Churu",
    "GMC Dungarpur": "GMC Dungarpur", "GMC Pali": "GMC Pali",
    "JMC Jhalawar": "JMC Jhalawar",
}


def clean(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def extract_2026() -> list[dict]:
    data = PDF_2026.read_bytes()
    assert hashlib.sha256(data).hexdigest() == PDF_2026_SHA256
    with fitz.open(stream=data, filetype="pdf") as pdf:
        assert len(pdf) == 2
        table = pdf[0].find_tables().tables[0].extract()
        assert "2026 State Counselling" in clean(table[0][2])
        # The attachment's second page is explicitly a 2025 private matrix.
        # Its numbers must never enter the 2026 comparison.
        assert "2025 State Counselling" in pdf[1].get_text()
        assert "Seats of Private Colleges" in pdf[1].get_text()
    assert len(table) == 38 and len(table[0]) == 31
    headers, subheaders = table[2], table[3]
    assert {clean(v) for v in headers if v} - {"S. No.", "Name of College"} == set(INSTITUTION_2026)
    assert all(clean(subheaders[i]) == "Total State Quota" for i in range(2, 9))
    assert all((clean(subheaders[i]), clean(subheaders[i + 1]), clean(subheaders[i + 2])) ==
               ("State (Govt)", "State (Mgmt)", "Total State Quota") for i in range(9, 30, 3))
    result = []
    branch_rows = [r for r in table if clean(r[0]).isdigit()]
    assert len(branch_rows) == 32 and [int(r[0]) for r in branch_rows] == list(range(1, 33))
    for row in branch_rows:
        branch = clean(row[1])
        values = [int(clean(v)) for v in row[2:31]]
        assert len(values) == 29 and sum(values[:7]) + sum(values[i] for i in range(9, 28, 3)) == values[-1]
        for i in range(9, 30, 3):
            assert values[i - 2] + values[i - 1] == values[i]
        for col in list(range(2, 9)) + list(range(9, 30, 3)):
            institution = INSTITUTION_2026[clean(headers[col])]
            seats = int(row[col]) if col < 9 else int(row[col]) + int(row[col + 1])
            result.append({"institution": institution, "branch": branch, "seats": seats})
    assert len(result) == 32 * 14 and sum(r["seats"] for r in result) == 990
    printed_totals = table[-1]
    assert clean(printed_totals[1]) == "Total Seats" and int(printed_totals[30]) == 990
    for col in list(range(2, 9)) + list(range(9, 30, 3)):
        institution = INSTITUTION_2026[clean(headers[col])]
        total = int(printed_totals[col]) if col < 9 else int(printed_totals[col]) + int(printed_totals[col + 1])
        assert total == sum(r["seats"] for r in result if r["institution"] == institution)
    return result


def build() -> None:
    source_2026 = extract_2026()
    by_year = {year: json.loads((ROOT / "data" / f"{year}-matrix.json").read_text(encoding="utf-8"))
               for year in (2024, 2025)}
    branches25 = {r["branch"] for r in by_year[2025] if r["round"] == "Round 1"}
    branches24 = {r["branch"] for r in by_year[2024] if r["round"] == "Round 1" and r["route"] == "MD/MS"}
    assert branches24 == set(BRANCH_2024)
    assert set(BRANCH_2024.values()) == branches25
    assert {r["branch"] for r in source_2026} == branches25

    totals: dict[int, dict[tuple[str, str], int]] = {2024: defaultdict(int), 2025: defaultdict(int), 2026: defaultdict(int)}
    source_pages: dict[int, dict[str, int]] = {2024: {}, 2025: {}}
    college_group: dict[str, str] = {}
    for year in (2024, 2025):
        for row in by_year[year]:
            if row["round"] != "Round 1" or row["route"] != "MD/MS":
                continue
            institution = row["institution"]
            branch = BRANCH_2024[row["branch"]] if year == 2024 else row["branch"]
            totals[year][(institution, branch)] += row["seats"]
            source_pages[year].setdefault(institution, row["pdf_page"])
            group = "Private" if row["institution_type"] == "Private" else "Government / RUHS / Society"
            assert institution not in college_group or college_group[institution] == group
            college_group[institution] = group
    for row in source_2026:
        key = row["institution"], row["branch"]
        assert college_group[key[0]] == "Government / RUHS / Society"
        totals[2026][key] += row["seats"]
    assert [sum(totals[y].values()) for y in (2024, 2025, 2026)] == [2008, 2378, 990]
    govt = [sum(v for (c, _), v in totals[y].items() if college_group[c] != "Private") for y in (2024, 2025, 2026)]
    assert govt == [945, 980, 990]

    institutions24 = {c for c, _ in totals[2024]}
    institutions25 = {c for c, _ in totals[2025]}
    institutions26 = {c for c, _ in totals[2026]}
    assert institutions24 <= institutions25 and institutions26 <= institutions25
    new_2025 = sum(v for (c, _), v in totals[2025].items() if c not in institutions24)
    assert new_2025 == 59  # ESIC Alwar 5 + Dr S S Tantia 54

    all_keys = set().union(*(set(t) for t in totals.values()))
    rows = []
    for institution, branch in sorted(all_keys):
        a = totals[2024].get((institution, branch), 0)
        b = totals[2025].get((institution, branch), 0)
        c = totals[2026].get((institution, branch), 0) if institution in institutions26 else None
        if not (a or b or c):
            continue
        source24 = f"{BASE_URL}documents/2024/matrix/round-1-medical.pdf#page={source_pages[2024][institution]}" if institution in institutions24 else ""
        source25 = f"{BASE_URL}documents/2025/matrix/round-1.pdf#page={source_pages[2025][institution]}"
        rows.append({
            "institution": institution, "group": college_group[institution], "branch": branch,
            "seats_2024": a, "seats_2025": b, "seats_2026": c,
            "change_2024_2025": b - a, "change_2025_2026": None if c is None else c - b,
            "college_status_2025": "New college" if institution not in institutions24 else "Existing college",
            "source_2024": source24, "source_2025": source25,
            "source_2026": f"{BASE_URL}documents/2026/matrix/round-1-medical.pdf#page=1" if c is not None else "",
        })
    assert [sum(r[f"seats_{y}"] or 0 for r in rows) for y in (2024, 2025, 2026)] == [2008, 2378, 990]
    assert sum(r["change_2024_2025"] for r in rows) == 370
    assert sum(r["change_2025_2026"] or 0 for r in rows if r["seats_2026"] is not None) == 10

    meta = {
        "basis": "Rajasthan Board Round 1 initially offered MD/MS seats; all quota labels summed within each college and branch",
        "totals_all_2024_2025": [2008, 2378],
        "totals_government_2024_2025_2026": govt,
        "new_college_seats_2025": new_2025,
        "existing_college_net_change_2025": 370 - new_2025,
        "source_2026": SOURCE_2026,
        "source_2026_sha256": PDF_2026_SHA256,
        "source_2026_scope": "Page 1 only: 2026 Government/RUHS/Society; page 2 is explicitly labelled 2025 private colleges and is excluded",
    }
    payload = {"meta": meta, "rows": rows}
    target = ROOT / "data" / "seat-growth-2024-2026.json"
    target.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    fields = list(rows[0])
    with (ROOT / "data" / "seat-growth-2024-2026.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

    wb = Workbook()
    notes = wb.active
    notes.title = "Read first"
    for note in [
        "Rajasthan Board Round 1 MD/MS initially offered seats, exam years 2024–2026.",
        "2024 and 2025 include government/RUHS/society and private colleges. 2026 includes government/RUHS/society only.",
        "The official 2026 attachment has a 2026 government page, but its private page is labelled 2025. Private 2026 is blank, not zero.",
        "All seat types within a college and branch are summed. These are Rajasthan state-counselling seats, not MCC AIQ intake.",
        "2024 short branch labels are mapped to the expanded 2025/2026 labels. The two 2024 anatomy and transfusion labels are combined.",
        "The 2026 government page is provisional, dated 16 October 2026. Later revisions may change seats.",
        "Each source column links to the original PDF page. A zero means no Round 1 seat in the verified source; blank 2026 private means unverified.",
        f"Official 2026 source: {SOURCE_2026}",
        f"2026 PDF SHA-256: {PDF_2026_SHA256}",
    ]:
        notes.append([note])
    notes.column_dimensions["A"].width = 120
    sheet = wb.create_sheet("Seat growth")
    labels = ["College", "Institution group", "Branch", "2024 seats", "2025 seats", "2026 seats",
              "2024 to 2025", "2025 to 2026", "2025 college status", "2024 source PDF", "2025 source PDF", "2026 source PDF"]
    sheet.append(labels)
    for row in rows:
        sheet.append([row[k] for k in fields])
        for col in (10, 11, 12):
            cell = sheet.cell(sheet.max_row, col)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0A6B69", underline="single")
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="184D55")
    for col, width in {"A": 55, "B": 28, "C": 55, "D": 15, "E": 15, "F": 15, "G": 18,
                       "H": 18, "I": 20, "J": 45, "K": 45, "L": 45}.items():
        sheet.column_dimensions[col].width = width
    table = Table(displayName="SeatGrowth", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    wb.save(ROOT / "data" / "Rajasthan-NEET-PG-2024-2026-seat-growth.xlsx")
    print(f"Verified Round 1 MD/MS growth: all 2024 2,008, 2025 2,378; government 945 -> 980 -> 990. {len(rows)} college/branch rows.")


if __name__ == "__main__":
    build()
